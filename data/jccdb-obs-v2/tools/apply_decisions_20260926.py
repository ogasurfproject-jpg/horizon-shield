# -*- coding: utf-8 -*-
"""
番人の判断(2026-09-26)を観測層 v2 に当てる。parser のあとに必ずこれを走らせる(何度走らせても同じ結果になる)。

1. 米国 SD / MT の DOT: 条件は「personal or informational use」まで(MT はさらに「改変しない限り」)。再配布の許可ではないので、
   値を写さない(published_restricted_not_copied、license restricted)。品目の並びは州が「copy or distribute ... for informational use」を
   許しているので残す。原本は raw_restricted/ へ。
2. 表そのものが複製・転載・電子媒体への加工を禁じている出典(中部 cbr、東北 thr、四国 skr、奈良県)は、値だけでなく行の一覧も
   公開の組み立てに入れない(observations_restricted/ へ移す)。件数は coverage に数として残す(数は事実)。
3. 建設工事費デフレーターの四半期別(2020年4-6月期以降)と年度別(2021年度以降)の 33 列は、月別に対して1列右にずれている疑いが強い
   (W2 の照合: 四半期の4期平均は年度別と一致し、月別とは範囲内で合わない。合成指数が個別指数の範囲に収まらない)。
   値を直さず、公開の組み立てから外して observations_hold/ に置く。国交省への照会待ち。
4. 近畿地整の台帳の license_quote が条文の要約だったので、条文そのものに置き換える(B1 が取り直した link.html)。
"""
import csv, json, os, re, shutil, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from obs_common import COLUMNS, write_obs, make_id

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
J = lambda *p: os.path.join(ROOT, *p)


def read(path):
    return list(csv.DictReader(open(path, encoding="utf-8")))


def ledger(sid):
    p = J("sources", sid + ".json")
    return p, json.load(open(p, encoding="utf-8"))


def save_ledger(p, d):
    json.dump(d, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


def us_restrict(csv_name, sid, reason):
    p = J("observations", "us", csv_name)
    if not os.path.exists(p):
        print("skip (無い):", csv_name); return
    rows = read(p)
    n = 0
    for r in rows:
        if r["price_status"] != "not_set":
            r["price_status"] = "published_restricted_not_copied"
        if r["price"]:
            n += 1
        r["price"] = ""
        r["ref_value"] = ""
        r["ref_note"] = ""
        r["license"] = "restricted"
        tag = "値は写していない: " + reason
        if tag not in r["note"]:
            r["note"] = (r["note"] + " / " if r["note"] else "") + tag
    write_obs(p, rows)
    lp, d = ledger(sid)
    d["license"] = "restricted"
    d["values_copied"] = False
    d["decision_20260926"] = reason
    save_ledger(lp, d)
    for ext in (".pdf", ".xlsx", ".csv"):
        src = J("raw", sid + ext)
        if os.path.exists(src):
            os.makedirs(J("raw_restricted"), exist_ok=True)
            shutil.move(src, J("raw_restricted", sid + ext))
    print("restricted:", csv_name, "値を外した行", n)


def move_listing(rel_csv, sid, reason):
    src = J("observations", rel_csv)
    dst = J("observations_restricted", rel_csv)
    if os.path.exists(src):
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.move(src, dst)
        print("行の一覧を公開の組み立てから外した:", rel_csv)
    lp, d = ledger(sid)
    d["row_listing"] = "not_in_public_build"
    d["row_listing_reason"] = reason
    save_ledger(lp, d)


def restore_nara_status():
    """奈良県の表は、状態(規格 x 地区ごとに「刊行物単価」「県が値を掲載」「空欄」のどれか)だけを公開の組み立てに戻す(2026-09-26 番人の判断、後半)。
    根拠: 奈良県サイトの条文は著作権の注記(「著作権法上認められた場合を除き、無断で複製・転用することはできません」)で、
    国の3局の表のような「磁気媒体入力を禁止」という著作権を超えた禁止ではない。規格の名前(JIS の呼び方)と、どの地区が刊行物単価かという事実は
    著作物ではない。値(県が掲載している金額)は写さない(県の表の中身そのもので、判断が分かれうるため)。"""
    src = J("observations_restricted", "jp", "material_nara_namacon_status_r8_09.csv")
    dst = J("observations", "jp", "material_nara_namacon_status_r8_09.csv")
    if os.path.exists(src):
        rows = read(src)
        assert all(r["price"] == "" for r in rows)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        shutil.move(src, dst)
        print("奈良県の表の状態を公開の組み立てに戻した:", len(rows))
    lp, d = ledger("nara-shizai-r8-09")
    if d.get("row_listing") != "status_only_listed":
        d["row_listing_previous"] = d.get("row_listing")
        d["row_listing"] = "status_only_listed"
        d["row_listing_reason"] = ("状態(刊行物単価 / 県が値を掲載 / 空欄)だけを載せ、値は写さない。奈良県サイトの条文は著作権の注記で、規格の名前と状態は事実であり著作物ではない。"
                                   "国の3局(中部・東北・四国)の表は『磁気媒体入力を禁止』と著作権を超えて禁じているので、そちらは件数だけ(2026-09-26 番人)。")
        save_ledger(lp, d)


def hold_deflator():
    moved = 0
    for fn in ("index_mlit_deflator.csv", "index_mlit_deflator_nendo.csv"):
        p = J("observations", "jp", fn)
        if not os.path.exists(p):
            continue
        rows = read(p)
        keep, hold = [], []
        for r in rows:
            (hold if "1列右にずれている疑い" in r["note"] else keep).append(r)
        if hold:
            write_obs(p, keep)
            hp = J("observations_hold", "jp", fn.replace(".csv", "_suspect_shift.csv"))
            prev = read(hp) if os.path.exists(hp) else []
            ids = {r["obs_id"] for r in prev}
            write_obs(hp, prev + [r for r in hold if r["obs_id"] not in ids])
            moved += len(hold)
    print("デフレーターの保留:", moved)


def realign_deflator():
    """保留したデフレーターの行を、月別と全セル照合したうえで正しい系列に付け直す(2026-09-26 番人の判断)。

    分かったこと(原本 deftsuki_2606.xlsx と defnendo_260630.xlsx を番人が直接読んで確かめた):
      - 四半期別 2020年4-6月期以降と年度別 2021年度以降で、見出し X の列にある値は、月別シートの「X の左隣の系列」の
        3か月平均(年度は 12か月平均)と一致する。四半期別 AT〜BY の 32列 x 25期 = 800 セル、年度別 AS〜BX の 32列 x 5年 = 160 セルが
        全部 0.1 以内で一致した。値は原本の数字のまま、系列の付け先だけを左隣に直す。
      - 最初の列(地方道路公社等の見出しの下)の値は、どの系列の平均とも合わない(四半期 22/25、年度 5/5 が不一致)。付け先が決まらないので使わない。
      - 最後の系列(S事務所・その他)は、その期間の値が原本のどの列にも無い(not_set)。
    付け直した行は observations/jp/index_mlit_deflator_realigned.csv。原本のとおりの行(見出しのとおりの付け先)は observations_hold/ に証拠として残す。
    """
    import openpyxl
    from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
    mfile = J("raw", "mlit-deflator-tsuki-2606.xlsx")
    if not os.path.exists(mfile):
        print("デフレーターの付け直し: 月別の原本が無いので何もしない"); return
    m = openpyxl.load_workbook(mfile, data_only=True).worksheets[0]
    months = {}
    for r in range(11, m.max_row + 1):
        y, mo = m.cell(r, 1).value, m.cell(r, 2).value
        if isinstance(y, int) and isinstance(mo, int):
            months[(y, mo)] = {c: m.cell(r, c).value for c in range(3, m.max_column + 1)}
    hdr = {c: m.cell(10, c).value for c in range(3, m.max_column + 1)}

    def mavg(keys, c):
        vals = [months.get(k, {}).get(c) for k in keys]
        if any(not isinstance(v, (int, float)) for v in vals):
            return None
        return sum(vals) / len(vals)

    out, report = [], {"quarterly": {}, "fiscal": {}}
    for fn, kind, offset in (("index_mlit_deflator_suspect_shift.csv", "quarterly", 0), ("index_mlit_deflator_nendo_suspect_shift.csv", "fiscal", 1)):
        hp = J("observations_hold", "jp", fn)
        if not os.path.exists(hp):
            continue
        rows = read(hp)
        byp = {}
        for r in rows:
            cell = r["source_page"].split("!")[1]
            col = CI(re.match(r"[A-Z]+", cell).group(0))
            byp.setdefault(r["period"], {})[col] = r
        ok = ng = unset = 0
        for period, cols in sorted(byp.items()):
            cs = sorted(cols)
            for i, col in enumerate(cs):
                h = cols[col]
                mcol = col + offset          # 月別シートの同じ系列の列(年度別は1列左から始まる)
                if hdr.get(mcol) != h["item_name"]:
                    raise SystemExit("見出しの対応が合わない: %s %s %s" % (fn, h["item_name"], hdr.get(mcol)))
                if kind == "quarterly":
                    y = int(period[:4]); qn = int(period[-1])
                    a = {2: 4, 3: 7, 4: 10, 1: 1}[qn]
                    keys = [(y, a), (y, a + 1), (y, a + 2)]
                else:
                    fy = int(period[2:])
                    keys = [(fy, mo) for mo in range(4, 13)] + [(fy + 1, mo) for mo in range(1, 4)]
                r2 = dict(h)
                r2["obs_id"] = make_id(h["source_id"], "realigned-20260926", h["item_name"], h["spec"], period)
                base = re.sub(r"注意\(W2 の照合\).*$", "", h["note"]).rstrip("。 ") + "。"
                if i + 1 < len(cs) and cs[i + 1] == col + 1:
                    src = cols[col + 1]
                    avg = mavg(keys, mcol)
                    v = float(src["price"])
                    if avg is None or abs(v - avg) > 0.1001:
                        raise SystemExit("照合に通らない: %s %s %s %s %s" % (fn, period, h["item_name"], v, avg))
                    r2["price"] = src["price"]
                    r2["price_status"] = "published_pdl"
                    r2["source_page"] = src["source_page"] + "(見出しは右隣の系列。値はこの系列)"
                    r2["note"] = base + ("原本では1列右の見出しの下にある値。月別シートのこの系列の%s平均と一致(差 %.2f、0.1 以内)を確かめて、この系列に付け直した(2026-09-26 番人。付け直した %d セルが全部一致)。"
                                         % ("3か月" if kind == "quarterly" else "12か月", abs(v - avg), 800 if kind == "quarterly" else 160))
                    ok += 1
                else:
                    r2["price"] = ""
                    r2["price_status"] = "not_set"
                    r2["source_page"] = ""
                    r2["note"] = base + "原本の%sでは、列ずれのため、この系列のこの期間の値がどの列にも無い(2026-09-26 番人)。月別シートの値はある。" % ("四半期別" if kind == "quarterly" else "年度別")
                    unset += 1
                out.append(r2)
        report[kind] = {"realigned": ok, "not_set": unset}
    # 最初の列(地方道路公社等)は、左隣の系列の値ではないので使わない。上の組み立てで、その列の値は付け先(左隣=阪神高速)を持たない。
    # 阪神高速(左隣)は範囲外で、自分の列に正しい値がある(W2 の照合)。
    write_obs(J("observations", "jp", "index_mlit_deflator_realigned.csv"), out)
    print("デフレーターの付け直し:", json.dumps(report, ensure_ascii=False))


def kkr_quote():
    lp, d = ledger("kkr-zairyo-r8-09")
    q = ("当ホームページで公開している情報（以下「コンテンツ」といいます。）の著作権は、特記されていない限り国土交通省近畿地方整備局に帰属し、"
         "権利表記の記載がない限り「公共データ利用規約（第1.0版）」（PDL1.0）に準拠した利用条件の下で、利用することができます。")
    if d.get("license_quote") != q:
        d["license_quote_previous"] = d.get("license_quote")
        d["license_quote"] = q
        d["license_url"] = "https://www.kkr.mlit.go.jp/link.html"
        d["license_checked"] = "2026-09-26 link.html を取り直し(last-modified Sun, 22 Mar 2026 11:48:04 GMT)。PDF の文字に権利表記・複製禁止の文言は無い(B1 が全頁で確認)。"
        save_ledger(lp, d)
        print("近畿の license_quote を条文に置き換えた")


# ---------------- 独立検証(V1 / V2、2026-09-26)を受けた直し ----------------
STAT_ZERO_BASES = {"count", "construction_cost_planned_total", "orders_received_total"}
OPEN_BY_LICENSE = {"PDL1.0": "published_pdl", "GOV-STD-2.0": "published_cc_by", "CC-BY-4.0": "published_cc_by",
                   "US-PD-17USC105": "public_domain", "OPEN-TERMS": "published_open_terms"}


def stat_zero():
    """統計の件数・金額の 0 は値(該当なし = 0)。not_set(設定していない)にしていた行を値 0 に戻す(V1 の指摘)。
    単価表の 0(市場性なし、設定なし)は not_set のまま。"""
    n = 0
    for fn in ("cost_sqft_mlit_chakko_2025.csv", "cost_sqft_mlit_chakko_2026_01_07.csv", "spending_mlit_reform_fy2025.csv"):
        p = J("observations", "jp", fn)
        if not os.path.exists(p):
            continue
        rows = read(p)
        ch = 0
        for r in rows:
            if r["price_status"] == "not_set" and r["price_basis"] in STAT_ZERO_BASES and (r["ref_value"] == "0" or "原本の値は 0" in r["note"]):
                why = r["ref_note"] or "原本の値は 0"
                r["price"] = "0"
                r["price_status"] = OPEN_BY_LICENSE[r["license"]]
                r["ref_value"] = ""
                r["ref_note"] = ""
                if why not in r["note"]:
                    r["note"] = (r["note"] + " / " if r["note"] else "") + why
                ch += 1
        if ch:
            write_obs(p, rows)
        n += ch
    print("統計の 0 を値に戻した行:", n)


def njdot_names():
    """NJDOT の品名 1 - 2" FLEXIBLE NONMETALLIC CONDUIT の『-』が parser で落ちていた(V2 の指摘、4 行)。原本の文字に戻す。"""
    p = J("observations", "us", "bid_item_njdot_2023q2.csv")
    if not os.path.exists(p):
        return
    rows = read(p)
    ch = 0
    for r in rows:
        if r["item_name"].startswith('1 2" FLEXIBLE NONMETALLIC CONDUIT'):
            r["item_name"] = r["item_name"].replace('1 2" FLEXIBLE', '1 - 2" FLEXIBLE', 1)
            ch += 1
    if ch:
        write_obs(p, rows)
    print("NJDOT の品名を直した行:", ch)


def sd_notes():
    """SD の note に残っていた原本の数(Bid Cnt、0.00 の量)を落とす。値を写さない出典なので、数は note にも残さない(V2 の指摘)。"""
    p = J("observations", "us", "bid_item_sddot_2024.csv")
    if not os.path.exists(p):
        return
    rows = read(p)
    ch = 0
    for r in rows:
        parts = [x for x in r["note"].split(" / ") if not re.search(r"Bid Cnt|原本の値は 0\.00|Total Quantity", x)]
        nn = " / ".join(parts)
        if nn != r["note"]:
            r["note"] = nn
            ch += 1
    if ch:
        write_obs(p, rows)
    print("SD の note から数を落とした行:", ch)


def ledger_quotes():
    """license_quote を条文そのままに戻す(V2 の指摘)。かぎ括弧を『』に置き換えていたもの、見出しと本文をつないでいたもの。"""
    import glob
    n = 0
    for f in glob.glob(J("sources", "*.json")):
        d = json.load(open(f, encoding="utf-8"))
        sid = d["source_id"]
        q = d.get("license_quote", "")
        q2 = q
        if sid.startswith("mlit-roumu-") or sid.startswith("mlit-gijutsusha-") or sid == "nara-shizai-r8-09":
            q2 = q2.replace("『", "「").replace("』", "」")
        if sid.startswith("estat-chakko-"):
            q2 = q2.replace("１）出典の記載について イ　", "１）出典の記載について / イ　").replace("６）その他 イ　", "６）その他 / イ　")
        if sid == "ogb-zairyo-r8-04" and " / 別のルールを適用するコンテンツについては" in q2:
            q2 = q2.split(" / 別のルールを適用するコンテンツについては")[0]
        if q2 != q:
            d.setdefault("license_quote_previous", q)
            d["license_quote"] = q2
            d["license_quote_fixed"] = "2026-09-26 独立検証(V2)で、条文の写し方(かぎ括弧の置き換え・見出しと本文の連結)が頁と文字どおりでないと分かり、頁の文字どおりに直した。利用条件の中身は変わらない。"
            json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            n += 1
    print("license_quote を条文どおりに直した台帳:", n)


def raw_copies():
    """原本が raw/ に無かった3出典を置く(V1 / V2 の指摘)。sha256 は台帳と照らしてから置く。"""
    import hashlib
    for sid, src, dst in (("mlit-roumu-r8", "/home/claude/work/src/mlit_roumu_r8_001981942.pdf", J("raw", "mlit-roumu-r8.pdf")),
                          ("nara-shizai-r8-09", "/home/claude/work/src/nara_r8_09_material.pdf", J("raw_restricted", "nara-shizai-r8-09.pdf")),
                          ("fhwa-nhcci", "/home/claude/work/jccdb_us/data/nhcci_r94d-n4f9_raw.json", J("raw", "fhwa-nhcci.json"))):
        if os.path.exists(dst) or not os.path.exists(src):
            continue
        d = json.load(open(J("sources", sid + ".json"), encoding="utf-8"))
        h = hashlib.sha256(open(src, "rb").read()).hexdigest()
        assert h == d["sha256"], (sid, h, d["sha256"])
        shutil.copy2(src, dst)
        print("原本を置いた:", dst)


# ---------------- 後半(2026-09-26 午後)の番人の判断 ----------------
US_ZERO_BASES = {"permit_valuation_usd", "permit_valuation_thousand_usd", "annual_total_thousand_usd",
                 "prevailing_fringe_hourly", "equipment_fuel_hourly", "equipment_fccm_hourly"}


def us_zero_values():
    """原本で 0 と書かれた統計の金額・付加給付・燃料費などは「0 という値」。not_set(設定なし)にしていた行を値 0 に戻す。
    単価の本体(時間あたり賃金の基本額、機械の時間単価の合計)の 0 は not_set のまま。都市の許可で有効な申告額が 1 件も無い区分(合計を出さない)も not_set のまま。"""
    import glob
    n = 0
    for p in sorted(glob.glob(J("observations", "us", "*.csv")) + glob.glob(J("observations_hold", "us", "*.csv"))):
        rows = read(p)
        ch = 0
        for r in rows:
            if r["price_status"] != "not_set" or r["price_basis"] not in US_ZERO_BASES:
                continue
            zero = r["ref_value"] in ("0", "0.0", "0.00") or re.search(r"原本の Value は 0|原本の値は 0|原本は \$0\.00|値 0\.00", r["note"] + " " + r["ref_note"])
            if not zero:
                continue
            r["price"] = "0"
            r["price_status"] = OPEN_BY_LICENSE[r["license"]]
            r["ref_value"] = ""
            r["ref_note"] = ""
            r["note"] = re.sub(r"(検査器が[^)]*not_set)", "(統計の 0 は値として持つ。2026-09-26 番人)", r["note"])
            r["note"] = re.sub(r"(単価の 0 は設定なしとし[^)]*)", "(0 は値として持つ。2026-09-26 番人)", r["note"])
            if "原本の値 0" not in r["note"]:
                r["note"] += " / 原本の値 0(該当なし・定めなし = 0)"
            ch += 1
        if ch:
            write_obs(p, rows)
        n += ch
    print("米国の統計・付加給付・燃料費の 0 を値に戻した行:", n)


def dbra_publish():
    """Davis-Bacon の賃金決定(公有の連邦の著作物)を公開の組み立てに入れる。
    取得の経路の記録: SAM.gov の頁の中から公開 API を呼んだ(鍵なし)。途中で SAM.gov の利用規約に自動取得の禁止があると分かり、担当は取得を止めた。
    内容は公有(17 U.S.C. 105、DOL の条文)なので、取れた 2,213 件は公開する。残りは自動取得せず、鍵のある公式の経路か情報公開請求で取る(2026-09-26 番人)。"""
    import glob
    moved = 0
    for src in sorted(glob.glob(J("observations_hold", "us", "labor_dol_davis_bacon_*.csv"))):
        dst = J("observations", "us", os.path.basename(src))
        shutil.move(src, dst)
        moved += 1
    fixed = 0
    for f in glob.glob(J("sources", "dol-dbra-*.json")):
        d = json.load(open(f, encoding="utf-8"))
        if "acquisition_note" not in d:
            d["acquisition_note"] = ("SAM.gov の頁を自動のブラウザで開き、頁の中から公開 API を呼んで取った。SAM.gov の利用規約は自動取得を禁じているので、"
                                     "気づいた時点で取得を止めた。値は連邦政府の著作物で公有(再配布は自由)。残りの決定は自動取得で取らない(2026-09-26 番人)。")
            json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
            fixed += 1
    print("Davis-Bacon を公開の組み立てに移したファイル:", moved, "/ 台帳に取得の記録を足した:", fixed)


def v5_fixes():
    """独立検証 V5(2026-09-26)の指摘を直す。値は1つも変えない(表記・台帳・名前だけ)。"""
    import glob, html
    out = {}
    # 1. HUD: HTML の文字参照が残っていた(COEUR D&apos;ALENE)
    p = J("observations", "us", "cost_limit_hud_tdc_2024.csv")
    if os.path.exists(p):
        rows = read(p); ch = 0
        for r in rows:
            for k in ("item_name", "spec", "area_label", "area_members", "geo_name", "note", "category"):
                v = html.unescape(r[k])
                if v != r[k]:
                    r[k] = v; ch += 1
        if ch: write_obs(p, rows)
        out["hud_unescape"] = ch
    # 2. DoD ACF: Site Name 欄が空の行に作文の表示を入れていた。原本の欄の空を note に書き、表示は原本の文字だけにする
    p = J("observations", "us", "index_dod_acf_2026.csv")
    if os.path.exists(p):
        rows = read(p); ch = 0
        for r in rows:
            if r["area_label"].startswith("(Site Name 空欄) "):
                r["area_label"] = r["area_label"][len("(Site Name 空欄) "):]
                r["note"] = (r["note"] + " / " if r["note"] else "") + "原本の Site Name 欄は空欄(area_label は Country と Location の欄の文字)"
                ch += 1
        if ch: write_obs(p, rows)
        out["dod_site_name"] = ch
    # 3. OEWS 都市圏: geo_name に州名を入れていた。都市圏の名前(原本の AREA_TITLE、area_label)にそろえる
    ch = 0
    for p in sorted(glob.glob(J("observations", "us", "wage_bls_oews_metro_*.csv"))):
        rows = read(p); c1 = 0
        for r in rows:
            if r["geo_level"] == "metro" and r["area_label"] and r["geo_name"] != r["area_label"]:
                r["geo_name"] = r["area_label"]; c1 += 1
        if c1: write_obs(p, rows)
        ch += c1
    out["oews_metro_geo_name"] = ch
    # 4. UFS の Industry Studies の行: 出所の詳細が原本に無いことを note に
    p = J("observations", "us", "cost_sqft_dod_ufc370101_2026.csv")
    if os.path.exists(p):
        rows = read(p); ch = 0
        for r in rows:
            if "Industry Stud" in (r["note"] + r["spec"] + r["ref_note"]) and "出所の詳細は原本に無い" not in r["note"]:
                r["note"] += " / 原本の出所欄は Industry Studies。どの調査かの出所の詳細は原本に無い(市販の物価資料の名は無いので値を入れた。2026-09-26 番人)"
                ch += 1
        if ch: write_obs(p, rows)
        out["ufs_industry_studies"] = ch
    # 5. FTA: 取得のときに失われた文字(U+FFFD)を '-' にした行に、仮の置き換えであることを書く
    p = J("observations", "us", "work_fta_capcost_2024_09.csv")
    if os.path.exists(p):
        rows = read(p); ch = 0
        for r in rows:
            if ("U+FFFD" in r["note"] or "置き換え" in r["note"]) and "仮の置き換え" not in r["note"]:
                r["note"] += " / 名前の中の '-' は、取得のときに失われた1文字の仮の置き換え(元の文字は分からない。値の列は影響なし)"
                ch += 1
        if ch: write_obs(p, rows)
        out["fta_placeholder"] = ch
    # 6. San Diego: ファイル名の期間を中身(2024〜2025)に合わせる
    for pre in ("spending", "cost_sqft"):
        a = J("observations", "us", "%s_city_permits_sandiego_2023_2025.csv" % pre)
        b = J("observations", "us", "%s_city_permits_sandiego_2024_2025.csv" % pre)
        if os.path.exists(a) and not os.path.exists(b):
            shutil.move(a, b); out.setdefault("renamed", []).append(os.path.basename(b))
    # 7. 台帳: NYC と San Diego の license_quote に、license_url 以外の頁の文が混ざっていた。分ける
    for f in glob.glob(J("sources", "city-nyc-*.json")):
        d = json.load(open(f, encoding="utf-8"))
        q = d["license_quote"]
        if "#Public Policies" in q and "license_quote_other" not in d:
            a, b = q.split("#Public Policies", 1)
            d["license_quote"] = a.strip()
            d["license_quote_other"] = {"where": "同じ Open Data Technical Standards Manual(GitHub CityOfNewYork/opendatatsm、commit fa204a27304262a68f83469c8468f7b00ca8f889)の Public Policies の節",
                                        "text": ("Public Policies" + b).strip()}
            json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=1); out.setdefault("ledger_split", []).append(d["source_id"])
    for f in glob.glob(J("sources", "city-sandiego-*.json")):
        d = json.load(open(f, encoding="utf-8"))
        q = d["license_quote"]
        if q.startswith("Dataset Details") and "license_quote_other" not in d:
            a, b = q.split("Terms of Use", 1)
            d["license_quote"] = ("Terms of Use" + b).strip()
            d["license_quote_other"] = {"where": d.get("landing") or "データセットの頁", "text": a.strip()}
            json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=1); out.setdefault("ledger_split", []).append(d["source_id"])
    # 8. Davis-Bacon の台帳の置き場所の記録を今に合わせる
    n = 0
    for f in glob.glob(J("sources", "dol-dbra-*.json")):
        d = json.load(open(f, encoding="utf-8"))
        if d.get("hold", "").startswith("observations_hold/"):
            d["hold_previous"] = d.pop("hold")
            d["placement"] = "observations/us/(2026-09-26 番人の判断で公開の組み立てに入れた。値は公有。取得の方法の記録は access_terms_note と acquisition_note)"
            json.dump(d, open(f, "w", encoding="utf-8"), ensure_ascii=False, indent=1); n += 1
    out["dbra_ledger_placement"] = n
    print("V5 の直し:", json.dumps(out, ensure_ascii=False))


def split_large(limit_mb=45):
    """GitHub の1ファイルの上限(警告 50MB、拒否 100MB)に合わせ、45MB を超える観測ファイルを行で2つ以上に分ける(_a, _b, ...)。"""
    import glob, math
    for p in sorted(glob.glob(J("observations", "**", "*.csv"), recursive=True)):
        size = os.path.getsize(p)
        if size <= limit_mb * 1000 * 1000:
            continue
        rows = read(p)
        k = math.ceil(size / (limit_mb * 1000 * 1000))
        per = math.ceil(len(rows) / k)
        base = p[:-4]
        for i in range(k):
            write_obs("%s_%s.csv" % (base, "abcdefghij"[i]), rows[i * per:(i + 1) * per])
        os.remove(p)
        print("大きいファイルを分けた:", os.path.basename(p), "->", k)


def main():
    us_restrict("bid_item_sddot_2024.csv", "sddot-bid-item-2024",
                "South Dakota の条件は『for personal or informational use』までで、再配布の許可ではない(2026-09-26 番人の判断)")
    us_restrict("bid_item_mtdot_2025.csv", "mtdot-wavg-2025",
                "Montana DOT の条件は『for personal or informational use』かつ『if the documents are not modified in any respect』で、観測層への並べ替えと再配布の許可ではない(2026-09-26 番人の判断)")
    move_listing("jp/material_cbr_zairyo_r8_10.csv", "cbr-zairyo-r8-10",
                 "原本 2頁: 本単価表を無断転載・複写や電子媒体等に加工することを禁じます。")
    move_listing("jp/material_thr_zairyo_r8_04.csv", "thr-zairyo-r8-04",
                 "原本 2頁: 本設計材料単価表の全部または一部を、無断で複製・転載・磁気媒体入力・販売することを禁止します。")
    move_listing("jp/material_skr_zairyo_r8_10.csv", "skr-zairyo-r8-10",
                 "原本 2頁: 本設計材料単価表の全部または一部を、第三者が複製・転載・磁気媒体入力・販売することを禁止します。")
    restore_nara_status()
    hold_deflator()
    realign_deflator()
    kkr_quote()
    stat_zero()
    njdot_names()
    sd_notes()
    ledger_quotes()
    raw_copies()
    dbra_publish()
    us_zero_values()
    v5_fixes()
    split_large()


if __name__ == "__main__":
    main()
