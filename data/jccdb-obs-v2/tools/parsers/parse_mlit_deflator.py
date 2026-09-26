# -*- coding: utf-8 -*-
"""
国土交通省「建設工事費デフレーター(2020年度基準)」を openpyxl でセルから直接読み、観測層 v2 の index 行にする。

入力: OBS2/raw/mlit-deflator-tsuki-2606.xlsx   (月次の公表ファイル deftsuki_2606.xlsx、令和8年8月31日付け)
        シート「deftsuki2020（月別）」2016年4月〜2026年6月、シート「deftsuki2020（四半期別）」
      OBS2/raw/mlit-deflator-nendo-260630.xlsx (年度次 defnendo_260630.xlsx、令和8年6月30日付け) 1951〜2025年度
出力: OBS2/observations/jp/index_mlit_deflator.csv        (月別 + 四半期別)
      OBS2/observations/jp/index_mlit_deflator_nendo.csv  (年度別)
      OBS2/reports/W2-deflator-checks.json

表頭は 3〜10 行目の 8 段。系列名は 10 行目(最下段)、spec に上の段からの経路(同じ語の続きは1つにまとめる)を原文で持つ。
値の無いセル(「***」、空)は not_set。
照合:
  Q 四半期の値 = その3か月の月別値の平均(小数第2位を四捨五入)か、全セル数える
  Y 年度の値 と 4〜3月の月別値の平均 の差(最大、件数)
  H 月別・四半期別・年度別の表頭 8 段が同じか
"""
import os, sys, json, hashlib, collections
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OBS2 = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(OBS2, "tools"))
from obs_common import make_id, write_obs  # noqa: E402

SID_M = "mlit-deflator-tsuki-2606"
SID_Y = "mlit-deflator-nendo-260630"
URL_M = "https://www.mlit.go.jp/statistics/details/content/deftsuki_2606.xlsx"
URL_Y = "https://www.mlit.go.jp/statistics/details/content/defnendo_260630.xlsx"
UNIT = "index (2020年度平均 = 100)"
QMAP = {"1-3月": 1, "4-6月": 2, "7-9月": 3, "10-12月": 4}
NOTE_BASE = "2020年度基準(2020年度平均=100)。2026年4月分の公表から2020年度基準に改定(原本注)。ウエイト項目に営業余剰や間接税(消費税を含む)等を含めない(原本注)"
NOTE_PROV = "2023年度以降は建設投資のウエイトに2020〜2022年度の平均値を暫定的に使用(原本注)"
SHIFT_FROM = 42  # 系列の番号(0始まり)。月別シートの列 AS「地方道路公社等」。ここから最後の「S事務所・その他」(BY)まで 33 系列
NOTE_SHIFT_Q = ("注意(W2 の照合): 2020年4-6月期以降の四半期別シートの「地方道路公社等」から「S事務所・その他」までの33列は、"
                "月別シートの同じ系列の3か月平均と合わず、左隣の系列の3か月平均と合う。四半期別の値が1列右にずれている疑い。値は原本のまま")
NOTE_SHIFT_Y = ("注意(W2 の照合): 2021年度以降の年度別ファイルの「地方道路公社等」から「S事務所・その他」までの33列は、"
                "月別シートの同じ系列の年度平均と合わず、左隣の系列の年度平均と合う。年度別の値が1列右にずれている疑い。値は原本のまま")


def col_letter(c):
    return openpyxl.utils.get_column_letter(c)


def headers(ws, c0, ncol):
    """列 c0..c0+ncol-1 の表頭(3〜10行)を (leaf, path) で返す"""
    out = []
    for c in range(c0, c0 + ncol):
        seq = [str(ws.cell(r, c).value) for r in range(3, 11)]
        path = []
        for s in seq:
            if not path or path[-1] != s:
                path.append(s)
        out.append((seq[-1], " > ".join(path), seq))
    # 表頭 8 段がまったく同じ列が2つある(「その他土木」の総合と、その内訳の「その他土木」)。列の位置で区別する
    cnt = collections.Counter(x[1] for x in out)
    out = [(lf, p + (" (原本の列 %s。同じ表頭の列が2つある)" % col_letter(c0 + i) if cnt[p] > 1 else ""), sq)
           for i, (lf, p, sq) in enumerate(out)]
    return out


def val(v):
    if isinstance(v, (int, float)):
        d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
        return ("num", d)
    if v is None:
        return ("blank", "")
    t = str(v).strip()
    if t in ("***", "", "-"):
        return ("blank", t)
    raise ValueError(repr(v))


def fmt(d):
    s = format(d, "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s


def main():
    wbm = openpyxl.load_workbook(os.path.join(OBS2, "raw", SID_M + ".xlsx"), data_only=True)
    wsm, wsq = wbm["deftsuki2020（月別）"], wbm["deftsuki2020（四半期別）"]
    wby = openpyxl.load_workbook(os.path.join(OBS2, "raw", SID_Y + ".xlsx"), data_only=True)
    wsy = wby["defnendo2020"]
    assert str(wsm.cell(1, 3).value) == "建設工事費デフレーター（2020年度基準）"
    assert str(wsm.cell(2, 3).value) == "［月別］" and str(wsq.cell(2, 3).value) == "［四半期別］"
    assert str(wsy.cell(2, 2).value) == "［年度別］"
    NS = 75
    hm = headers(wsm, 3, NS)
    hq = headers(wsq, 3, NS)
    hy = headers(wsy, 2, NS)
    checks = collections.OrderedDict()
    checks["H_month_vs_quarter_same"] = [x[2] for x in hm] == [x[2] for x in hq]
    checks["H_month_vs_year_same"] = [x[2] for x in hm] == [x[2] for x in hy]
    assert checks["H_month_vs_quarter_same"] and checks["H_month_vs_year_same"]
    assert len({x[1] for x in hm}) == NS, "経路が一意でない"
    assert str(wsm.cell(10, NS + 2).value) == "S事務所・その他" and wsm.cell(10, NS + 3).value is None
    rows_m, rows_y = [], []
    monthly = {}  # (series_idx, 'YYYY-MM') -> Decimal
    # ---- 月別 ----
    r = 11
    while True:
        y, m = wsm.cell(r, 1).value, wsm.cell(r, 2).value
        if not isinstance(y, int):
            break
        vals = [val(wsm.cell(r, 3 + i).value) for i in range(NS)]
        if all(v[0] == "blank" for v in vals):
            r += 1
            continue  # 未公表の月(空欄の行)
        per = "%04d-%02d" % (y, m)
        fy = y if m >= 4 else y - 1
        for i, (leaf, path, _) in enumerate(hm):
            kind, v = vals[i]
            notes = [NOTE_BASE] + ([NOTE_PROV] if fy >= 2023 else [])
            row = dict(obs_id=make_id(SID_M, "月別", path, per), country="JP", layer="index",
                       category="建設工事費デフレーター（2020年度基準） 月別", item_name=leaf, spec=path, unit=UNIT,
                       geo_level="national", geo_code="JP", geo_name="日本", currency="", price_basis="index_value",
                       period=per, source_id=SID_M, source_page="deftsuki2020（月別）!%s%d" % (col_letter(3 + i), r),
                       evidence_url=URL_M, license="PDL1.0")
            if kind == "num":
                row.update(price=fmt(v), price_status="published_pdl")
                monthly[(i, per)] = v
            else:
                row.update(price="", price_status="not_set")
                notes.append("原本は「%s」" % v if v else "原本は空欄")
            row["note"] = "。".join(notes)
            rows_m.append(row)
        r += 1
    checks["month_first_last"] = [rows_m[0]["period"], rows_m[-1]["period"]]
    # ---- 四半期別 ----
    qn = collections.Counter()
    qmax = Decimal(0)
    r = 11
    while True:
        y, ql = wsq.cell(r, 1).value, wsq.cell(r, 2).value
        if not isinstance(y, int):
            break
        vals = [val(wsq.cell(r, 3 + i).value) for i in range(NS)]
        if all(v[0] == "blank" for v in vals):
            r += 1
            continue
        q = QMAP[str(ql)]
        per = "%04dQ%d" % (y, q)
        months = ["%04d-%02d" % (y, 3 * (q - 1) + k) for k in (1, 2, 3)]
        fy = y if q >= 2 else y - 1
        for i, (leaf, path, _) in enumerate(hq):
            kind, v = vals[i]
            notes = [NOTE_BASE, "原本の期間表示「%d年 %s」(暦年の四半期)" % (y, ql)] + ([NOTE_PROV] if fy >= 2023 else [])
            row = dict(obs_id=make_id(SID_M, "四半期別", path, per), country="JP", layer="index",
                       category="建設工事費デフレーター（2020年度基準） 四半期別", item_name=leaf, spec=path, unit=UNIT,
                       geo_level="national", geo_code="JP", geo_name="日本", currency="", price_basis="index_value",
                       period=per, source_id=SID_M, source_page="deftsuki2020（四半期別）!%s%d" % (col_letter(3 + i), r),
                       evidence_url=URL_M, license="PDL1.0")
            region = "late_tail" if (y, q) >= (2020, 2) and i >= SHIFT_FROM else "other"
            if region == "late_tail":
                notes.append(NOTE_SHIFT_Q)
            if kind == "num":
                row.update(price=fmt(v), price_status="published_pdl")
                ms = [monthly.get((i, mm)) for mm in months]
                ml = [monthly.get((i - 1, mm)) for mm in months] if i > 0 else [None]
                if all(x is not None for x in ms):
                    avg = (sum(ms) / 3).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                    d = abs(avg - v)
                    same = d <= Decimal("0.1")
                    left = all(x is not None for x in ml) and abs((sum(ml) / 3).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) - v) <= Decimal("0.1")
                    qn["%s:%s" % (region, "same" if same else ("left_only" if left else "neither"))] += 1
                    if region == "other":
                        qmax = max(qmax, d)
                else:
                    qn["Q_nomonth"] += 1
            else:
                row.update(price="", price_status="not_set")
                notes.append("原本は「%s」" % v if v else "原本は空欄")
            row["note"] = "。".join(notes)
            rows_m.append(row)
        r += 1
    checks["Q"] = dict(qn)
    checks["Q_absdiff_max_other"] = fmt(qmax)
    # ---- 年度別 ----
    yn = collections.Counter()
    ymax = Decimal(0)
    ydiffs = []
    r = 11
    while True:
        lab = wsy.cell(r, 1).value
        if lab is None:
            break
        s = str(lab)
        fy = int(s[:4])
        prov = "暫定" in s
        per = "FY%04d" % fy
        for i, (leaf, path, _) in enumerate(hy):
            kind, v = val(wsy.cell(r, 2 + i).value)
            notes = [NOTE_BASE, "原本の年度表示「%s」" % s] + ([NOTE_PROV] if prov or fy >= 2023 else [])
            row = dict(obs_id=make_id(SID_Y, "年度別", path, per), country="JP", layer="index",
                       category="建設工事費デフレーター（2020年度基準） 年度別", item_name=leaf, spec=path, unit=UNIT,
                       geo_level="national", geo_code="JP", geo_name="日本", currency="", price_basis="index_value",
                       period=per, source_id=SID_Y, source_page="defnendo2020!%s%d" % (col_letter(2 + i), r),
                       evidence_url=URL_Y, license="PDL1.0")
            region = "late_tail" if fy >= 2021 and i >= SHIFT_FROM else "other"
            if region == "late_tail":
                notes.append(NOTE_SHIFT_Y)
            if kind == "num":
                row.update(price=fmt(v), price_status="published_pdl")
                mlist = [(fy if mm >= 4 else fy + 1, mm) for mm in (4, 5, 6, 7, 8, 9, 10, 11, 12, 1, 2, 3)]
                ms = [monthly.get((i, "%04d-%02d" % t)) for t in mlist]
                ml = [monthly.get((i - 1, "%04d-%02d" % t)) for t in mlist] if i > 0 else [None]
                if all(x is not None for x in ms):
                    avg = (sum(ms) / 12).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                    d = abs(avg - v)
                    same = d <= Decimal("0.1")
                    left = all(x is not None for x in ml) and abs((sum(ml) / 12).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP) - v) <= Decimal("0.1")
                    yn["%s:%s" % (region, "same" if same else ("left_only" if left else "neither"))] += 1
                    if region == "other" and d > ymax:
                        ymax = d
                    if region == "other" and d > Decimal("0.1"):
                        ydiffs.append([per, path, fmt(v), fmt(avg)])
            else:
                row.update(price="", price_status="not_set")
                notes.append("原本は「%s」" % v if v else "原本は空欄")
            row["note"] = "。".join(notes)
            rows_y.append(row)
        r += 1
    checks["Y"] = dict(yn)
    checks["Y_absdiff_max_other"] = fmt(ymax)
    checks["Y_diffs_gt_0.1_other"] = ydiffs[:20]
    files = {}
    for fn, rows in (("index_mlit_deflator.csv", rows_m), ("index_mlit_deflator_nendo.csv", rows_y)):
        p = os.path.join(OBS2, "observations", "jp", fn)
        n = write_obs(p, rows)
        files["observations/jp/" + fn] = {
            "rows": n, "by_status": dict(collections.Counter(x["price_status"] for x in rows)),
            "periods": len({x["period"] for x in rows}), "series": len({x["spec"] for x in rows}),
            "sha256": hashlib.sha256(open(p, "rb").read()).hexdigest()}
    rep = {"files": files, "checks": checks}
    json.dump(rep, open(os.path.join(OBS2, "reports", "W2-deflator-checks.json"), "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    print(json.dumps(rep, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
