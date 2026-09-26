# -*- coding: utf-8 -*-
"""
U1-us-bls: BLS 職業別雇用・賃金統計(OEWS)May 2025 の建設関係の職業を観測層 v2 に入れる。

入力(OBS2/raw、2026-09-26 に Apify web-fetch formats=raw で取得):
  - bls-oesm25nat.zip  https://www.bls.gov/oes/special-requests/oesm25nat.zip (national_M2025_dl.xlsx)
  - bls-oesm25st.zip   https://www.bls.gov/oes/special-requests/oesm25st.zip  (state_M2025_dl.xlsx)
      州の zip は Apify のデータ項目の上限(9.4MB)を超えたため、Range ヘッダで 0-3999999 と 4000000- の2回に分けて取り、
      つないだ。2回とも同じ ETag、合計 bytes が Content-Range の全長 7563829 と一致、zip の CRC 検査が通ることを確かめた。
出力:
  - observations/us/wage_bls_oews_2025.csv
  - sources/bls-oews-m2025-national.json, sources/bls-oews-m2025-state.json
  - reports/U1-us-bls_oews_emp_check.csv(全国の雇用者数と州の合計)
セルは openpyxl で直接読む。対象の職業: SOC 47-xxxx の全行(major, minor, broad, detailed)と、下の EXTRA。
記号 "*" "**" "#" は値を入れず not_set、note に記号の意味を Field Descriptions の Notes の原文で。
"""
import csv, collections, hashlib, io, json, os, sys, zipfile
HERE = os.path.dirname(os.path.abspath(__file__))
OBS2 = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, os.path.join(OBS2, "tools"))
from obs_common import make_id, num, write_obs, US_STATES
import openpyxl

RAW = os.path.join(OBS2, "raw")
RETRIEVED = "2026-09-26"
PERIOD = "2025-05"
EXTRA = ["11-9021", "13-1051", "17-1011", "17-2051", "17-3011", "17-3022", "49-9021"]
SRC = collections.OrderedDict([
    ("bls-oews-m2025-national", ("bls-oesm25nat.zip", "oesm25nat/national_M2025_dl.xlsx",
                                 "https://www.bls.gov/oes/special-requests/oesm25nat.zip", "Fri, 15 May 2026 13:12:23 GMT")),
    ("bls-oews-m2025-state", ("bls-oesm25st.zip", "oesm25st/state_M2025_dl.xlsx",
                              "https://www.bls.gov/oes/special-requests/oesm25st.zip", "Fri, 15 May 2026 13:19:24 GMT")),
])
# (列名, price_basis, 単位, PRSE の列)
METRICS = [
    ("TOT_EMP", "count", "persons (employment)", "EMP_PRSE"),
    ("H_MEAN", "wage_hourly_mean", "USD/hour", "MEAN_PRSE"),
    ("A_MEAN", "wage_annual_mean", "USD/year", "MEAN_PRSE"),
    ("H_PCT10", "wage_hourly_p10", "USD/hour", None),
    ("H_PCT25", "wage_hourly_p25", "USD/hour", None),
    ("H_MEDIAN", "wage_hourly_median", "USD/hour", None),
    ("H_PCT75", "wage_hourly_p75", "USD/hour", None),
    ("H_PCT90", "wage_hourly_p90", "USD/hour", None),
    ("A_PCT10", "wage_annual_p10", "USD/year", None),
    ("A_PCT25", "wage_annual_p25", "USD/year", None),
    ("A_MEDIAN", "wage_annual_median", "USD/year", None),
    ("A_PCT75", "wage_annual_p75", "USD/year", None),
    ("A_PCT90", "wage_annual_p90", "USD/year", None),
]
LICENSE_QUOTE = ("The Bureau of Labor Statistics (BLS) is a Federal government agency and everything that we publish, "
                 "both in hard copy and electronically, is in the public domain, except for previously copyrighted photographs "
                 "and illustrations. You are free to use our public domain material without specific permission, although we do "
                 "ask that you cite the Bureau of Labor Statistics as the source.")


def sha(path):
    b = open(path, "rb").read()
    return hashlib.sha256(b).hexdigest(), len(b)


def load(zfn, member):
    z = zipfile.ZipFile(os.path.join(RAW, zfn))
    assert z.testzip() is None
    wb = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    notes = {}
    for r in wb["Field Descriptions"].iter_rows(values_only=True):
        t = r[0]
        if isinstance(t, str):
            for sym in ("**", "*", "#", "~"):
                if t.startswith(sym + " ") and "=" in t and sym not in notes:
                    notes[sym] = t.strip()
                    break
    return ws.title, rows, notes


def main():
    out_rows, stat = [], collections.Counter()
    emp = {}                       # (src, area, code, o_group) -> TOT_EMP
    info = {}
    wage_check = collections.Counter()
    for src, (zfn, member, url, lm) in SRC.items():
        sheet, rows, notes = load(zfn, member)
        H = list(rows[0]); ix = {h: i for i, h in enumerate(H)}
        majors = {r[ix["OCC_CODE"]]: r[ix["OCC_TITLE"]] for r in rows[1:] if r[ix["O_GROUP"]] == "major"}
        n_sel = 0
        for ln, r in enumerate(rows[1:], start=2):
            code = r[ix["OCC_CODE"]]
            if not (code.startswith("47-") or code in EXTRA):
                continue
            assert r[ix["NAICS"]] == "000000" and r[ix["OWN_CODE"]] == "1235"
            n_sel += 1
            area, og = r[ix["AREA"]], r[ix["O_GROUP"]]
            if r[ix["AREA_TYPE"]] == "1":
                geo = dict(geo_level="national", geo_code="US", geo_name="United States")
            else:
                geo = dict(geo_level="state", geo_code=area, geo_name=US_STATES[area])
            base = dict(country="US", layer="wage", category=majors[code[:2] + "-0000"], item_name=r[ix["OCC_TITLE"]],
                        spec="SOC %s; O_GROUP %s; cross-industry (NAICS 000000); ownership 1235" % (code, og),
                        area_label=r[ix["AREA_TITLE"]], area_code=area, period=PERIOD, source_id=src,
                        source_page="%s row %d" % (sheet, ln), evidence_url=url, license="US-PD-17USC105", **geo)
            vals = {}
            for col, pb, unit, prse in METRICS:
                v = r[ix[col]]
                row = dict(base, obs_id=make_id(src, area, code, og, col), unit=unit, price_basis=pb,
                           currency="" if pb == "count" else "USD")
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    row.update(price=num(repr(v) if isinstance(v, float) else str(v)), price_status="public_domain")
                    vals[col] = v
                    if prse:
                        pv = r[ix[prse]]
                        if isinstance(pv, (int, float)):
                            row.update(ref_value=num(repr(pv) if isinstance(pv, float) else str(pv)),
                                       ref_note="%s (percent relative standard error, 原本の列 %s)" % (prse, prse))
                    stat["public_domain"] += 1
                else:
                    s = "" if v is None else str(v).strip()
                    if s in notes:
                        row.update(price="", price_status="not_set", note='OEWS の記号 "%s": %s' % (s, notes[s]))
                        stat["not_set:" + s] += 1
                    elif s == "":
                        row.update(price="", price_status="not_set", note="原本のセルが空欄")
                        stat["not_set:blank"] += 1
                    else:
                        raise SystemExit("想定外のセル: %s %s %s %r" % (src, ln, col, v))
                out_rows.append(row)
            emp[(src, area, code, og)] = vals.get("TOT_EMP")
            info[(code, og)] = r[ix["OCC_TITLE"]]
            # 別の読み方: 年収 = 時給 x 2080 か(OEWS の年収は時給 x 2080 で作る。年収は 10 ドル単位に丸め)
            for h, a in (("H_MEAN", "A_MEAN"), ("H_PCT10", "A_PCT10"), ("H_PCT25", "A_PCT25"), ("H_MEDIAN", "A_MEDIAN"),
                         ("H_PCT75", "A_PCT75"), ("H_PCT90", "A_PCT90")):
                if h in vals and a in vals:
                    d = abs(vals[h] * 2080 - vals[a])
                    wage_check["pairs"] += 1
                    wage_check["within_rounding_bound_15.4usd"] += d <= 15.4   # 時給 0.005 x 2080 + 年収 5
                    wage_check["max_diff_x100"] = max(wage_check["max_diff_x100"], int(round(d * 100)))
        stat["selected_rows:" + src] = n_sel
        stat["file_rows:" + src] = len(rows) - 1

    out = os.path.join(OBS2, "observations", "us", "wage_bls_oews_2025.csv")
    n = write_obs(out, out_rows)

    # 照合: 全国の雇用者数 と 州の合計(51 = 50州+DC、+ 領土3)
    nat = "bls-oews-m2025-national"; st = "bls-oews-m2025-state"
    terr = {"66", "72", "78"}
    chk = []
    for (src, area, code, og), v in emp.items():
        if src != nat:
            continue
        s51 = s54 = 0; sup = 0; nst = 0
        for (s2, a2, c2, o2), v2 in emp.items():
            if s2 == st and c2 == code and o2 == og:
                nst += 1
                if v2 is None:
                    sup += 1
                    continue
                s54 += v2
                if a2 not in terr:
                    s51 += v2
        if nst == 0:
            s51 = s54 = None
        chk.append(dict(occ_code=code, o_group=og, title=info[(code, og)], national_tot_emp=v, states_reporting=nst,
                        states_suppressed_emp=sup, sum_50states_dc=s51, sum_with_territories=s54,
                        ratio_50states_dc=("%.4f" % (s51 / v)) if (v and s51 is not None) else "",
                        ratio_with_territories=("%.4f" % (s54 / v)) if (v and s54 is not None) else "",
                        note="" if nst else "州のファイルにこの階層(%s)の行が無い" % og))
    ck = os.path.join(OBS2, "reports", "U1-us-bls_oews_emp_check.csv")
    with open(ck, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(chk[0].keys()), lineterminator="\n")
        w.writeheader(); w.writerows(chk)

    # 照合: 階層の合計。全国: 47-0000 = minor 5 群の合計、47-0000 = detailed の合計。
    # 州: 州のファイルには major と detailed しか無い(minor, broad が無い)ので 47-0000 = detailed の合計(雇用者数が ** の行があれば不完全)
    hier = collections.Counter()
    for src in SRC:
        areas = sorted({a for (s, a, c, o) in emp if s == src})
        for a in areas:
            maj = emp.get((src, a, "47-0000", "major"))
            det = [v for (s, a2, c, o), v in emp.items() if s == src and a2 == a and c.startswith("47-") and o == "detailed"]
            if src == nat:
                mins = [emp.get((src, a, c, "minor")) for c in ("47-1000", "47-2000", "47-3000", "47-4000", "47-5000")]
                hier["national_major_minus_sum_minor"] = maj - sum(mins)
                hier["national_major_minus_sum_detailed"] = maj - sum(det)
                continue
            # 州では公表されない職業の行そのものが無いことがあるので、47-0000 >= detailed(値のある行)の合計 の下限の照合
            if maj is None:
                hier["state_major_suppressed"] += 1
                continue
            d = maj - sum(v for v in det if v is not None)
            hier["state_areas"] += 1
            hier["state_major_ge_sum_detailed_minus_50"] += d >= -50
            hier["state_min_major_minus_sum"] = min(hier.get("state_min_major_minus_sum", d), d)
            hier["state_areas_with_suppressed_detailed"] += any(v is None for v in det)

    # 台帳
    for src, (zfn, member, url, lm) in SRC.items():
        h, b = sha(os.path.join(RAW, zfn))
        led = collections.OrderedDict([
            ("source_id", src), ("country", "US"),
            ("title", "Occupational Employment and Wage Statistics (OEWS), May 2025, %s (%s)" %
             ("national" if "national" in src else "state", member.split("/")[-1])),
            ("publisher", "U.S. Bureau of Labor Statistics"), ("url", url), ("landing", "https://www.bls.gov/oes/tables.htm"),
            ("retrieved_at", RETRIEVED), ("bytes", b), ("sha256", h), ("http_last_modified", lm),
            ("license", "US-PD-17USC105"), ("license_url", "https://www.bls.gov/opub/copyright-information.htm"),
            ("license_quote", LICENSE_QUOTE),
            ("attribution", "Source: U.S. Bureau of Labor Statistics, Occupational Employment and Wage Statistics, May 2025 (https://www.bls.gov/oes/)"),
            ("how_read", "zip の中の xlsx を openpyxl でセルごとに読む(%s の 1 枚目)。対象は SOC 47-xxxx の全行と %s、"
                         "cross-industry(NAICS 000000)・全所有形態(OWN_CODE 1235)。1 行の職業から TOT_EMP と時給・年収の平均と分位"
                         "(10/25/50/75/90)の 13 観測を作る。記号 * ** # は値を入れず not_set、note に Field Descriptions の Notes の原文。"
                         "O_GROUP が broad と detailed の両方で出る職業(OEWS の仕様)は両方を残し、spec に O_GROUP を書く。"
                         "照合: 全国の TOT_EMP と州の合計(reports/U1-us-bls_oews_emp_check.csv)、47-0000 と minor 5 群の合計、"
                         "年収 = 時給 x 2080 の全組。" % (member.split("/")[-1], ", ".join(EXTRA))),
            ("values_copied", True), ("fetched_via", "Apify web-fetch (formats=raw)" +
                                      ("; Range ヘッダで 2 回に分けて取得し連結(同じ ETag、合計 bytes = Content-Range の全長、zip CRC 検査済み)" if "state" in src else "")),
            ("published", "2026-05"), ("scope_quote", "May 2025 OEWS Estimates"),
            ("rows", sum(1 for r in out_rows if r["source_id"] == src)),
        ])
        json.dump(led, open(os.path.join(OBS2, "sources", src + ".json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    rat = [float(c["ratio_50states_dc"]) for c in chk if c["states_reporting"]]
    print(json.dumps(dict(rows=n, stat=dict(stat), wage_x2080=dict(wage_check), hierarchy=dict(hier),
                          emp_check=dict(occupations=len(chk), ratio_min=min(rat), ratio_max=max(rat),
                                         compared=len(rat), not_in_state_file=sum(1 for c in chk if not c["states_reporting"]),
                                         within_0_1pct=sum(1 for x in rat if abs(x - 1) <= 0.001),
                                         within_1pct=sum(1 for x in rat if abs(x - 1) <= 0.01),
                                         with_suppressed=sum(1 for c in chk if c["states_suppressed_emp"]),
                                         no_suppression_ratio_range=[min(float(c["ratio_50states_dc"]) for c in chk if c["states_reporting"] and not c["states_suppressed_emp"]),
                                                                     max(float(c["ratio_50states_dc"]) for c in chk if c["states_reporting"] and not c["states_suppressed_emp"])])),
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
